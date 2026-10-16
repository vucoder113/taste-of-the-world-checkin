import io
import json
import unittest
from unittest.mock import MagicMock, patch

from openpyxl import Workbook

import app as report_app


def make_workbook():
    workbook = Workbook()
    target = workbook.active
    target.title = "Ds_target_nv"
    target.append(["Tên nhân viên", "Chức vụ", "Phòng ban", "Số khách target"])
    target.append(["Nguyễn An", "Quản lý", "Kinh doanh", 10])

    registered = workbook.create_sheet("Ds_kh_đk_qr")
    registered.append(["Mã KH", "Họ tên khách hàng", "Điện thoại", "Người mời", "VIP"])
    registered.append(["KH01", "Khách Một", "0900000001", "Nguyễn An", "VIP"])
    registered.append(["KH02", "Khách Hai", "0900000002", "Nguyễn An", "Không"])
    registered.append(["KH03", "Khách Ba", "0900000003", "Nguyễn An", "K VIP"])
    registered.append(["KH04", "Khách Bốn", "0900000004", "Trần Bình", "Không"])

    afternoon = workbook.create_sheet("Ds_kh_checkin_chieu")
    afternoon.append(["Mã KH", "Họ tên khách hàng", "Người mời"])
    afternoon.append(["KH01", "Khách Một", "Nguyễn An"])

    dinner = workbook.create_sheet("Ds_kh_checkin_tiectoi")
    dinner.append(["Mã KH", "Họ tên khách hàng", "Người mời"])
    dinner.append(["KH01", "Khách Một", "Nguyễn An"])

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def make_separate_files():
    workbooks = []

    registered = Workbook()
    sheet = registered.active
    sheet.title = "Sheet1"
    sheet.append(["DANH SÁCH ĐĂNG KÝ"])
    sheet.append(["STT", "Ho ten", "SDT", "Email", "CÔNG TY", "CHỨC VỤ", "THÀNH PHỐ", "BỘ PHẬN MỜI", "NHÂN VIÊN MỜI"])
    sheet.append([1, "Khách Một", "0900000001", "one@example.com", "Công ty Một", "Giám đốc", "HN", "Kinh Doanh", "Nguyễn An"])
    sheet.append([2, "Khách Hai", "0900000002", "two@example.com", "Công ty Hai", "Quản lý", "HN", "Kinh Doanh", "Nguyễn An"])
    workbooks.append(("registrations.xlsx", registered))

    afternoon = Workbook()
    sheet = afternoon.active
    sheet.append(["DANH SÁCH KHÁCH CHECK-IN BUỔI CHIỀU"])
    sheet.append(["STT", "Họ & tên", "SDT", "Email", "Công ty", "Chức vụ", "Người mời", "Vip"])
    sheet.append([1, "Khách Một", "0900000001", "one@example.com", "Công ty Một", "Giám đốc", "Anh Nguyễn An", "ü"])
    sheet.append([2, "Khách Hai", "0999999999", "other@example.com", "Công ty Khác", "Khách", "Tên không có trong target", None])
    workbooks.append(("afternoon_checkin.xlsx", afternoon))

    dinner = Workbook()
    sheet = dinner.active
    sheet.title = "Ds_kh_checkin_tiectoi"
    sheet.append(["DANH SÁCH KHÁCH MỜI TIỆC TỐI"])
    sheet.append(["STT", "Họ và Tên Khách mời", "SL Khách", "Vip", "Sale", "Check in"])
    sheet.append([1, "Khách Một", 2, "ü", "Anh Nguyễn An", "ü"])
    sheet.append([2, "Nhóm khách Hai", 3, "Không", "Nguyễn An", None])
    sheet.append([3, "Nhóm khách Chưa rõ SL", None, "ü", "Nguyễn An", "ü"])
    workbooks.append(("dinner_guests.xlsx", dinner))

    employees = Workbook()
    sheet = employees.active
    sheet.append(["STT", "Tên nhân viên", "Chức vụ", "Phòng ban", "Số khách target"])
    sheet.append([1, "Nguyễn An", "Quản lý", "Kinh doanh", 8])
    sheet.append([2, "Bùi Bình", "Trưởng nhóm", "BOD", 5])
    sheet.append([3, "An Bình", "Nhân viên", "BOD", 4])
    workbooks.append(("targets.xlsx", employees))

    result = []
    for filename, workbook in workbooks:
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        result.append((filename, output.getvalue()))
    return result


def make_hcm_files():
    registered = Workbook()
    sheet = registered.active
    sheet.append(["STT", "Ho ten", "SDT", "Email", "Thanh pho", "Nhan vien moi"])
    sheet.append([1, "Khách HCM", "0900000001", "hcm@example.com", "HCM", "Nguyễn An"])
    sheet.append([2, "Khách HN", "0900000002", "hn@example.com", "Hà Nội", "Nguyễn An"])

    afternoon = Workbook()
    sheet = afternoon.active
    sheet.append(["STT", "Ho ten", "SDT", "Thanh pho", "Nhân viên mời"])
    sheet.append([1, "Khách chiều HCM", "0900000003", "TP. Hồ Chí Minh", "Nguyễn An"])
    sheet.append([2, "Khách chiều HN", "0900000004", "HN", "Nguyễn An"])

    dinner = Workbook()
    sheet = dinner.active
    sheet.append(["DANH SÁCH KHÁCH MỜI TIỆC TỐI"])
    sheet.append(["STT", "Kênh", "Tên đại diện", "Tên công ty", "SL khách tham dự", "Số thăm", "Check-in"])
    sheet.append(["KÊNH CH", None, None, None, None, None, None])
    sheet.append([1, "Kênh CH", "Đại diện Một", "Công ty Một", 1, "45", "✓"])
    sheet.append([2, "Supplier & VIP", "Đại diện Hai", "Công ty Hai", 2, None, "☐"])

    employees = Workbook()
    sheet = employees.active
    sheet.append(["Tên nhân viên", "Chức vụ", "Phòng ban", "Số khách target"])
    sheet.append(["Nguyễn An", "Quản lý", "Kinh doanh", 8])

    result = []
    for filename, workbook in (
        ("DS KHÁCH ĐĂNG KÍ QR.xlsx", registered),
        ("DS KHÁCH CHECK IN CHIỀU.xlsx", afternoon),
        ("DS_Khach_Moi_Tiec_Toi.xlsx", dinner),
        ("DS TARGET KHÁCH.xlsx", employees),
    ):
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        result.append((filename, output.getvalue()))
    return result


class ReportAppTests(unittest.TestCase):
    def setUp(self):
        report_app.app.config["TESTING"] = True
        self.client = report_app.app.test_client()
        report_app._current_reports = {}

    def test_upload_builds_linked_employee_report(self):
        response = self.client.post(
            "/api/upload",
            data={"file": (io.BytesIO(make_workbook()), "guests.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 200)
        report = response.get_json()["report"]
        self.assertEqual(report["totals"], {"employees": 2, "registered": 4, "afternoon": 1, "dinner": 1})
        employee = report["employees"][0]
        self.assertEqual(employee["name"], "Nguyễn An")
        self.assertEqual((employee["target"], employee["invited"], employee["afternoon"], employee["dinner"]), ("10", 3, 1, 1))
        self.assertAlmostEqual(employee["attendance_rate"], 2 / 10)
        self.assertEqual(len(employee["guests"]), 3)
        guest = next(item for item in employee["guests"] if item["customer_id"] == "KH01")
        self.assertTrue(guest["vip"])
        self.assertTrue(guest["registered"] and guest["afternoon"] and guest["dinner"])
        non_vip_guest = next(item for item in employee["guests"] if item["customer_id"] == "KH03")
        self.assertFalse(non_vip_guest["vip"])
        self.assertEqual(report["charts"]["dinner"][0]["vip"], 1)
        self.assertEqual(report["charts"]["registered"][0]["name"], "Nguyễn An")
        self.assertEqual(report["charts"]["registered"][0]["count"], 3)

    def test_recognizes_common_inviter_header_variants(self):
        workbook = Workbook()
        target = workbook.active
        target.title = "Ds_target_nv"
        target.append(["Tên nhân viên", "Chức vụ", "Phòng ban", "Target"])
        target.append(["Nguyễn An", "Quản lý", "Kinh doanh", 5])
        registered = workbook.create_sheet("Ds_kh_dk_qr")
        registered.append(["Mã KH", "Tên khách", "NV phụ trách"])
        registered.append(["KH01", "Khách Một", "Nguyễn An"])
        afternoon = workbook.create_sheet("Ds_kh_checkin_chieu")
        afternoon.append(["Mã KH", "Tên khách", "NV phụ trách"])
        dinner = workbook.create_sheet("Ds_kh_checkin_tiectoi")
        dinner.append(["Mã KH", "Tên khách", "NV phụ trách"])
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()

        report = report_app.parse_workbook(output.getvalue())
        self.assertEqual(report["charts"]["registered"][0]["name"], "Nguyễn An")

    def test_groups_missing_inviter_as_unassigned(self):
        workbook = Workbook()
        target = workbook.active
        target.title = "Ds_target_nv"
        target.append(["Tên nhân viên"])
        target.append(["Nguyễn An"])
        registered = workbook.create_sheet("Ds_kh_dk_qr")
        registered.append(["Mã KH", "Tên khách"])
        registered.append(["KH01", "Khách Một"])
        afternoon = workbook.create_sheet("Ds_kh_checkin_chieu")
        afternoon.append(["Mã KH", "Tên khách"])
        dinner = workbook.create_sheet("Ds_kh_checkin_tiectoi")
        dinner.append(["Mã KH", "Tên khách"])
        dinner.append(["KH01", "Khách Một"])
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()

        report = report_app.parse_workbook(output.getvalue())
        self.assertEqual(report["charts"]["dinner"], [
            {"name": "Chưa xác định", "vip": 0, "non_vip": 1, "count": 1},
        ])

    def test_uses_sale_column_as_dinner_inviter(self):
        workbook = Workbook()
        target = workbook.active
        target.title = "Ds_target_nv"
        target.append(["Tên nhân viên"])
        target.append(["Nguyễn An"])
        registered = workbook.create_sheet("Ds_kh_dk_qr")
        registered.append(["Mã KH", "Tên khách"])
        registered.append(["KH01", "Khách Một"])
        afternoon = workbook.create_sheet("Ds_kh_checkin_chieu")
        afternoon.append(["Mã KH", "Tên khách"])
        dinner = workbook.create_sheet("Ds_kh_checkin_tiectoi")
        dinner.append(["Mã KH", "Tên khách", "Sale"])
        dinner.append(["KH01", "Khách Một", "Nguyễn An"])
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()

        report = report_app.parse_workbook(output.getvalue())
        self.assertEqual(report["charts"]["dinner"], [
            {"name": "Nguyễn An", "vip": 0, "non_vip": 1, "count": 1},
        ])

    def test_rejects_workbook_missing_required_sheet(self):
        workbook = Workbook()
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        response = self.client.post(
            "/api/upload",
            data={"file": (io.BytesIO(output.getvalue()), "wrong.xlsx")},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("Thiếu sheet bắt buộc", response.get_json()["message"])

    def test_report_endpoint_requires_upload(self):
        response = self.client.get("/api/report")
        self.assertEqual(response.status_code, 404)

    def test_report_endpoint_rejects_unknown_region(self):
        report_app._current_reports = {"HN": {"totals": {}}}

        response = self.client.get("/api/report?region=unknown")

        self.assertEqual(response.status_code, 400)

    def test_report_endpoint_reads_persisted_report(self):
        expected = {"totals": {"employees": 1}}
        with (
            patch.object(report_app, "DATABASE_URL", "postgresql://test"),
            patch.object(report_app, "read_persisted_report", return_value=expected) as read_report,
        ):
            response = self.client.get("/api/report")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["report"], expected)
        self.assertEqual(response.get_json()["reports"], {"HN": expected})
        read_report.assert_called_once_with()

    def test_persisted_report_storage_uses_private_schema_and_tls(self):
        report = {"guest": "Khách Một"}
        connection = MagicMock()
        connection_context = MagicMock()
        connection_context.__enter__.return_value = connection
        with (
            patch.object(report_app, "DATABASE_URL", "postgresql://test"),
            patch.object(report_app.psycopg, "connect", return_value=connection_context) as connect,
        ):
            report_app.write_persisted_report(report)

        connect.assert_called_once_with(
            "postgresql://test", connect_timeout=10, sslmode="require"
        )
        statements = [call.args[0] for call in connection.execute.call_args_list]
        self.assertTrue(any("CREATE SCHEMA IF NOT EXISTS app_private" in sql for sql in statements))
        self.assertTrue(any("ENABLE ROW LEVEL SECURITY" in sql for sql in statements))
        self.assertTrue(any("REVOKE ALL ON SCHEMA app_private FROM PUBLIC" in sql for sql in statements))
        self.assertEqual(json.loads(connection.execute.call_args.args[1][0]), report)

    def test_upload_persists_report_before_returning_success(self):
        with (
            patch.object(report_app, "DATABASE_URL", "postgresql://test"),
            patch.object(report_app, "read_persisted_report", return_value=None),
            patch.object(report_app, "write_persisted_report") as write_report,
        ):
            response = self.client.post(
                "/api/upload",
                data={"file": (io.BytesIO(make_workbook()), "guests.xlsx")},
                content_type="multipart/form-data",
            )

        self.assertEqual(response.status_code, 200)
        write_report.assert_called_once_with({"regions": {"HN": response.get_json()["report"]}})

    def test_persisted_hcm_upload_preserves_hanoi_report(self):
        hanoi = {"totals": {"employees": 10}}
        with (
            patch.object(report_app, "DATABASE_URL", "postgresql://test"),
            patch.object(report_app, "read_persisted_report", return_value={"regions": {"HN": hanoi}}),
            patch.object(report_app, "write_persisted_report") as write_report,
        ):
            response = self.client.post(
                "/api/upload",
                data={
                    "region": "HCM",
                    "files": [(io.BytesIO(content), filename) for filename, content in make_hcm_files()],
                },
                content_type="multipart/form-data",
            )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            write_report.call_args.args[0],
            {"regions": {"HN": hanoi, "HCM": response.get_json()["report"]}},
        )

    def test_upload_reports_persistence_failure(self):
        previous = {"totals": {"employees": 1}}
        report_app._current_reports = {"HN": previous}
        with (
            patch.object(report_app, "DATABASE_URL", "postgresql://test"),
            patch.object(report_app, "read_persisted_report", return_value={"regions": {"HN": previous}}),
            patch.object(
                report_app,
                "write_persisted_report",
                side_effect=report_app.psycopg.OperationalError("database unavailable"),
            ),
        ):
            response = self.client.post(
                "/api/upload",
                data={"file": (io.BytesIO(make_workbook()), "guests.xlsx")},
                content_type="multipart/form-data",
            )

        self.assertEqual(response.status_code, 503)
        self.assertEqual(report_app._current_reports, {"HN": previous})

    def test_authentication_protects_dashboard_and_report_api(self):
        with (
            patch.object(report_app, "AUTH_REQUIRED", True),
            patch.object(report_app, "APP_PASSWORD", "test-password"),
            patch.dict(report_app.app.config, {"SECRET_KEY": "test-secret"}),
        ):
            response = self.client.get("/")
            self.assertEqual(response.status_code, 302)
            self.assertIn("/login", response.headers["Location"])
            api_response = self.client.get("/api/report")
            self.assertEqual(api_response.status_code, 401)
            invalid_login = self.client.post("/login", data={"password": "wrong-password"})
            self.assertEqual(invalid_login.status_code, 401)
            login_response = self.client.post("/login", data={"password": "test-password"})
            self.assertEqual(login_response.status_code, 302)
            dashboard_response = self.client.get("/")
            self.assertEqual(dashboard_response.status_code, 200)
            self.assertIn("Đăng xuất", dashboard_response.get_data(as_text=True))
            logout_response = self.client.post("/logout")
            self.assertEqual(logout_response.status_code, 302)
            self.assertEqual(self.client.get("/api/report").status_code, 401)

    def test_dashboard_page_renders(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertIn("Khách mời &amp; check-in", response.get_data(as_text=True))
        self.assertIn("CHECK-IN BUỔI CHIỀU", response.get_data(as_text=True))
        self.assertIn("TỶ LỆ THAM DỰ", response.get_data(as_text=True))
        self.assertIn("Chưa đối chiếu Target", response.get_data(as_text=True))
        self.assertIn("dinner-adjustment-form", response.get_data(as_text=True))

    def test_reads_four_separate_workbooks_without_merging_lists(self):
        report = report_app.parse_separate_files(make_separate_files())
        self.assertEqual(report["totals"], {
            "employees": 3,
            "registered": 2,
            "afternoon": 2,
            "dinner_invited": 5,
            "dinner": 2,
            "dinner_unquantified": 1,
            "dinner_unquantified_checkins": 1,
        })
        employee = next(item for item in report["employees"] if item["name"] == "Nguyễn An")
        self.assertEqual(
            (employee["target"], employee["invited"], employee["afternoon"], employee["dinner_invited"], employee["dinner"]),
            ("8", 2, 1, 5, 2),
        )
        self.assertEqual(employee["attendance_rate"], 3 / 8)
        self.assertEqual(len(employee["guests"]), 6)
        self.assertEqual(
            [item["name"] for item in report["employees"]],
            ["An Bình", "Bùi Bình", "Nguyễn An"],
        )
        self.assertEqual(report["employees"][0]["attendance_rate"], 0)
        self.assertEqual(report["totals"]["afternoon"], 2)
        self.assertEqual(report["charts"]["afternoon"][0]["count"], 1)
        self.assertEqual(sum(item["count"] for item in report["charts"]["afternoon"]), 2)
        self.assertIn(
            "Tên không có trong target",
            [item["name"] for item in report["charts"]["afternoon"]],
        )
        self.assertEqual(report["unmatched"], {
            "registered": 0,
            "afternoon": 1,
            "dinner_invited": 0,
            "dinner": 0,
        })
        registered_guest = next(
            guest for guest in employee["guests"]
            if guest["source"] == "registered" and guest["name"] == "Khách Một"
        )
        self.assertEqual(
            (registered_guest["registered"], registered_guest["afternoon"], registered_guest["dinner_invited"], registered_guest["dinner"]),
            (True, True, True, True),
        )
        guest_two = next(
            guest for guest in employee["guests"]
            if guest["source"] == "registered" and guest["name"] == "Khách Hai"
        )
        self.assertFalse(guest_two["afternoon"])
        unmatched_guest = next(
            guest for guest in employee["guests"]
            if guest["source"] == "registered" and guest["name"] == "Khách Hai"
        )
        self.assertFalse(unmatched_guest["afternoon"])
        self.assertEqual(report["charts"]["registered"][0]["unclassified"], 2)
        self.assertEqual(report["charts"]["dinner_invited"][0]["count"], 5)
        self.assertEqual(report["charts"]["dinner"][0]["count"], 2)
        self.assertTrue(any("1 dòng tiệc tối" in warning for warning in report["warnings"]))
        self.assertIsNone(employee["guests"][-1]["quantity"])
        self.assertEqual(
            next(source for source in report["sources"] if source["category"] == "dinner")["guests"],
            5,
        )

    def test_resolves_unique_employee_aliases_but_not_ambiguous_names(self):
        employees = [
            {"name": "Đào Văn Tâm", "position": "NV KD", "department": "Kinh Doanh"},
            {"name": "CHÚ HƯNG", "position": "HĐQT", "department": "BOD"},
            {"name": "Nguyễn Tiến Hưng", "position": "NV KD MT", "department": "Kinh Doanh"},
            {"name": "ANH TÙNG", "position": "TGĐ", "department": "BOD"},
        ]
        self.assertEqual(
            report_app.resolve_target_employee("Anh Tâm", employees)["name"],
            "Đào Văn Tâm",
        )
        self.assertEqual(
            report_app.resolve_target_employee("TGĐ", employees)["name"],
            "ANH TÙNG",
        )
        self.assertEqual(
            report_app.resolve_target_employee("TDG", employees)["name"],
            "ANH TÙNG",
        )
        self.assertEqual(
            report_app.resolve_target_employee("Sếp tùng", employees)["name"],
            "ANH TÙNG",
        )
        self.assertIsNone(report_app.resolve_target_employee("Anh Hưng", employees))

    def test_attendance_rate_uses_registration_count_when_it_meets_target(self):
        employees = [
            {
                "name": "Registered above target",
                "department": "",
                "invited": 12,
                "target": "10",
                "afternoon": 2,
                "dinner": 1,
            },
            {
                "name": "Registered at target",
                "department": "",
                "invited": 10,
                "target": "10",
                "afternoon": 2,
                "dinner": 1,
            },
        ]
        prepared = report_app.prepare_employee_report(employees)
        self.assertEqual(prepared[0]["attendance_rate"], 3 / 12)
        self.assertEqual(prepared[1]["attendance_rate"], 3 / 10)

    def test_upload_accepts_the_four_separate_workbooks(self):
        response = self.client.post(
            "/api/upload",
            data={"files": [(io.BytesIO(content), filename) for filename, content in make_separate_files()]},
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["report"]["totals"]["dinner"], 2)

    def test_hcm_parser_filters_city_rows_and_reads_dinner_without_inviter(self):
        report = report_app.parse_separate_files(make_hcm_files(), region="HCM")

        self.assertEqual(report["totals"], {
            "employees": 1,
            "registered": 1,
            "afternoon": 1,
            "dinner_invited": 3,
            "dinner": 1,
            "dinner_unquantified": 0,
            "dinner_unquantified_checkins": 0,
        })
        self.assertEqual(report["chart_group_by"]["dinner"], "Kênh")
        self.assertEqual(
            [(item["name"], item["count"]) for item in report["charts"]["dinner"]],
            [("Kênh CH", 1)],
        )
        self.assertEqual(
            [(item["name"], item["count"]) for item in report["charts"]["dinner_invited"]],
            [("Supplier & VIP", 2), ("Kênh CH", 1)],
        )
        self.assertEqual(report["unmatched"]["dinner_invited"], 3)
        self.assertTrue(any("miền khác" in warning for warning in report["warnings"]))
        self.assertTrue(any("không có cột người mời" in warning for warning in report["warnings"]))
        self.assertFalse(any("dòng khách mời tiệc tối chưa có tên người mời" in warning for warning in report["warnings"]))

    def test_manual_hcm_dinner_checkin_adjustment_updates_unattributed_chart_only(self):
        report = report_app.parse_separate_files(make_hcm_files(), region="HCM")
        employee_checkins = sum(employee["dinner"] for employee in report["employees"])

        report_app.set_dinner_checkin_total(report, 7, "Điều chỉnh theo yêu cầu người dùng")

        self.assertEqual(report["totals"]["dinner"], 7)
        self.assertEqual(sum(group["count"] for group in report["charts"]["dinner"]), 7)
        self.assertEqual(
            next(group for group in report["charts"]["dinner"] if group["name"] == "Điều chỉnh thủ công")["count"],
            6,
        )
        self.assertEqual(report["unmatched"]["dinner"], 7)
        self.assertEqual(sum(employee["dinner"] for employee in report["employees"]), employee_checkins)
        self.assertEqual(report["manual_adjustments"]["dinner_checkins"]["base_total"], 1)
        self.assertTrue(any("7 khách (tăng 6" in warning for warning in report["warnings"]))

        report_app.set_dinner_checkin_total(report, 7, "Điều chỉnh theo yêu cầu người dùng")
        self.assertEqual(sum(group["count"] for group in report["charts"]["dinner"]), 7)

        report_app.set_dinner_checkin_total(report, 1, "Hủy điều chỉnh")
        self.assertEqual(report["totals"]["dinner"], 1)
        self.assertFalse(any(group["name"] == "Điều chỉnh thủ công" for group in report["charts"]["dinner"]))
        self.assertNotIn("manual_adjustments", report)

    def test_manual_hcm_adjustment_api_updates_current_report(self):
        report_app._current_reports = {
            "HCM": report_app.parse_separate_files(make_hcm_files(), region="HCM")
        }

        response = self.client.post(
            "/api/report/adjust",
            json={
                "region": "HCM",
                "dinner_checkins": 120,
                "reason": "Điều chỉnh theo yêu cầu người dùng",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["report"]["totals"]["dinner"], 120)
        persisted = self.client.get("/api/report?region=HCM").get_json()["report"]
        self.assertEqual(persisted["totals"]["dinner"], 120)

    def test_manual_hcm_adjustment_api_persists_to_supabase(self):
        report = report_app.parse_separate_files(make_hcm_files(), region="HCM")
        stored_reports = {"regions": {"HN": {"totals": {"dinner": 10}}, "HCM": report}}
        with (
            patch.object(report_app, "DATABASE_URL", "postgresql://test"),
            patch.object(report_app, "read_persisted_report", return_value=stored_reports),
            patch.object(report_app, "write_persisted_report") as write_report,
        ):
            response = self.client.post(
                "/api/report/adjust",
                json={
                    "region": "HCM",
                    "dinner_checkins": 7,
                    "reason": "Dashboard test adjustment",
                },
            )

        self.assertEqual(response.status_code, 200)
        persisted_payload = write_report.call_args.args[0]
        self.assertEqual(persisted_payload["regions"]["HCM"]["totals"]["dinner"], 7)
        self.assertEqual(persisted_payload["regions"]["HN"], {"totals": {"dinner": 10}})

    def test_uploading_hcm_preserves_hanoi_report(self):
        hanoi_response = self.client.post(
            "/api/upload",
            data={
                "region": "HN",
                "files": [(io.BytesIO(content), filename) for filename, content in make_separate_files()],
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(hanoi_response.status_code, 200)

        hcm_response = self.client.post(
            "/api/upload",
            data={
                "region": "HCM",
                "files": [(io.BytesIO(content), filename) for filename, content in make_hcm_files()],
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(hcm_response.status_code, 200)
        self.assertEqual(set(hcm_response.get_json()["reports"]), {"HN", "HCM"})

        restored = self.client.get("/api/report").get_json()
        self.assertEqual(set(restored["reports"]), {"HN", "HCM"})
        self.assertEqual(restored["reports"]["HN"], hanoi_response.get_json()["report"])
        self.assertEqual(restored["reports"]["HCM"], hcm_response.get_json()["report"])


if __name__ == "__main__":
    unittest.main()
