from __future__ import annotations

import re
import hmac
import json
import threading
import unicodedata
from collections import defaultdict
from io import BytesIO
from os import environ
from pathlib import Path
from typing import Any

from flask import Flask, jsonify, redirect, render_template, request, session, url_for
from openpyxl import load_workbook
import psycopg
from werkzeug.exceptions import RequestEntityTooLarge


BASE_DIR = Path(__file__).resolve().parent
REQUIRED_SHEETS = {
    "dskhcheckintiectoi": "dinner",
    "dskhcheckinchieu": "afternoon",
    "dskhdkqr": "registered",
    "dstargetnv": "employees",
}
SHEET_LABELS = {
    "dinner": "Check-in tiệc tối",
    "afternoon": "Check-in chiều",
    "registered": "Đăng ký QR",
}
ALIASES = {
    "employee": (
        "ten nhan vien", "ho ten nhan vien", "nhan vien", "ten nv", "ho ten nv",
        "ten sale", "sale", "ten nguoi moi", "nguoi moi", "sale phu trach",
        "nhan vien phu trach",
    ),
    "position": ("chuc vu", "vi tri", "position"),
    "department": ("phong ban", "bo phan", "department"),
    "target": ("so khach target", "khach target", "target", "chi tieu", "so luong target"),
    "customer_id": ("ma khach hang", "ma kh", "ma khach", "customer id", "customer code"),
    "phone": ("so dien thoai", "dien thoai", "sdt", "phone", "mobile"),
    "email": ("email", "e mail", "thu dien tu", "email address"),
    "customer_name": ("ho va ten", "ho ten khach hang", "ten khach hang", "ten kh", "ten khach", "ten dai dien", "ho ten", "name"),
    "inviter": (
        "sale", "ten sale", "sale moi", "sale moi khach", "sale phu trach",
        "ten nguoi moi", "nguoi moi", "nhan vien moi", "nv moi",
        "nhan vien moi khach", "nv moi khach", "nguoi moi khach",
        "nhan vien phu trach", "nv phu trach",
        "nguoi phu trach", "ten nguoi phu trach", "ten sale", "sale",
        "nhan vien", "ten nhan vien", "nguoi gioi thieu",
        "nhan vien gioi thieu", "nv gioi thieu",
    ),
    "vip": ("vip", "loai vip", "phan loai vip", "hang khach hang", "loai khach", "phan loai", "phan hang", "nhom khach"),
    "guest_count": ("sl khach", "so khach", "so luong khach", "guest count", "quantity"),
    "checked_in": ("check in", "checkin", "da check in", "check in status"),
    "inviter_department": ("bo phan moi", "phong ban moi", "don vi moi"),
    "city": ("thanh pho", "city", "tinh thanh", "province"),
    "channel": ("kenh", "channel"),
}

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024
AUTH_REQUIRED = environ.get("REQUIRE_AUTH", "").casefold() in {"1", "true", "yes"}
APP_PASSWORD = environ.get("APP_PASSWORD", "")
SECRET_KEY = environ.get("SECRET_KEY", "")
DATABASE_URL = environ.get("DATABASE_URL", "")
if AUTH_REQUIRED and (not APP_PASSWORD or not SECRET_KEY):
    raise RuntimeError("REQUIRE_AUTH cần được bật cùng APP_PASSWORD và SECRET_KEY.")
if AUTH_REQUIRED and not DATABASE_URL:
    raise RuntimeError("Website triển khai cần DATABASE_URL để lưu báo cáo bền vững.")
app.config.update(
    SECRET_KEY=SECRET_KEY or None,
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=AUTH_REQUIRED,
)
_report_lock = threading.Lock()
_current_reports: dict[str, dict[str, Any]] = {}
REGIONS = {"HN": "Hà Nội", "HCM": "TP. Hồ Chí Minh"}


def _ensure_report_table(connection: psycopg.Connection) -> None:
    connection.execute("CREATE SCHEMA IF NOT EXISTS app_private")
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS app_private.reports (
            id smallint PRIMARY KEY CHECK (id = 1),
            payload jsonb NOT NULL,
            updated_at timestamptz NOT NULL DEFAULT now()
        )
        """
    )
    connection.execute("ALTER TABLE app_private.reports ENABLE ROW LEVEL SECURITY")
    connection.execute("REVOKE ALL ON SCHEMA app_private FROM PUBLIC")
    connection.execute("REVOKE ALL ON TABLE app_private.reports FROM PUBLIC")


def read_persisted_report() -> dict[str, Any] | None:
    with psycopg.connect(DATABASE_URL, connect_timeout=10, sslmode="require") as connection:
        _ensure_report_table(connection)
        row = connection.execute(
            "SELECT payload FROM app_private.reports WHERE id = 1"
        ).fetchone()
    return row[0] if row else None


def reports_from_payload(payload: dict[str, Any] | None) -> dict[str, dict[str, Any]]:
    if not payload:
        return {}
    regions = payload.get("regions")
    if isinstance(regions, dict):
        return {
            region: report
            for region, report in regions.items()
            if region in REGIONS and isinstance(report, dict)
        }
    return {"HN": payload}


def set_dinner_checkin_total(report: dict[str, Any], total: int, reason: str) -> dict[str, Any]:
    if total < 0:
        raise ValueError("Số khách check-in tiệc tối không được âm.")

    adjustments = report.get("manual_adjustments", {})
    previous = adjustments.get("dinner_checkins", {})
    base_total = previous.get("base_total", report["totals"]["dinner"])
    delta = total - base_total
    if delta < 0:
        raise ValueError(
            f"Số check-in điều chỉnh ({total}) thấp hơn số đã ghi nhận trong dữ liệu nguồn ({base_total})."
        )

    charts = report.setdefault("charts", {})
    dinner_chart = charts.setdefault("dinner", [])
    manual_group = next(
        (group for group in dinner_chart if group.get("name") == "Điều chỉnh thủ công"),
        None,
    )
    if delta:
        if manual_group is None:
            manual_group = {
                "name": "Điều chỉnh thủ công",
                "vip": 0,
                "non_vip": 0,
                "unclassified": delta,
                "count": delta,
            }
            dinner_chart.append(manual_group)
        else:
            manual_group.update({"vip": 0, "non_vip": 0, "unclassified": delta, "count": delta})
    elif manual_group is not None:
        dinner_chart.remove(manual_group)

    report["totals"]["dinner"] = total
    if "unmatched" in report:
        source_unmatched = previous.get("base_unmatched", report["unmatched"].get("dinner", 0))
        report["unmatched"]["dinner"] = source_unmatched + delta

    if delta:
        adjustments = report.setdefault("manual_adjustments", {})
        adjustments["dinner_checkins"] = {
            "base_total": base_total,
            "base_unmatched": previous.get("base_unmatched", report.get("unmatched", {}).get("dinner", 0) - delta),
            "total": total,
            "delta": delta,
            "reason": reason,
        }
    else:
        adjustments.pop("dinner_checkins", None)
        if not adjustments:
            report.pop("manual_adjustments", None)

    message = (
        f"Điều chỉnh thủ công check-in tiệc tối: {total} khách "
        f"(tăng {delta} so với dữ liệu nguồn). Lý do: {reason}"
    )
    warnings = report.setdefault("warnings", [])
    warnings[:] = [warning for warning in warnings if not warning.startswith("Điều chỉnh thủ công check-in tiệc tối:")]
    if delta:
        warnings.append(message)
    return report


def write_persisted_report(report: dict[str, Any]) -> None:
    payload = json.dumps(report, ensure_ascii=False)
    with psycopg.connect(DATABASE_URL, connect_timeout=10, sslmode="require") as connection:
        _ensure_report_table(connection)
        connection.execute(
            """
            INSERT INTO app_private.reports (id, payload)
            VALUES (1, %s::jsonb)
            ON CONFLICT (id) DO UPDATE
            SET payload = EXCLUDED.payload, updated_at = now()
            """,
            (payload,),
        )


@app.before_request
def require_login():
    if not AUTH_REQUIRED or request.endpoint in {"login", "health", "static"}:
        return None
    if session.get("authenticated") is True:
        return None
    if request.path.startswith("/api/"):
        return jsonify({"ok": False, "message": "Vui lòng đăng nhập để tiếp tục."}), 401
    return redirect(url_for("login"))


@app.get("/health")
def health():
    return jsonify({"ok": True})


@app.route("/login", methods=["GET", "POST"])
def login():
    if not AUTH_REQUIRED:
        return redirect(url_for("index"))
    error = None
    if request.method == "POST":
        password = request.form.get("password", "")
        if hmac.compare_digest(password, APP_PASSWORD):
            session.clear()
            session["authenticated"] = True
            return redirect(url_for("index"))
        error = "Mật khẩu chưa đúng. Vui lòng thử lại."
    return render_template("login.html", error=error), 401 if error else 200


@app.post("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


def normalize(value: Any) -> str:
    text = str(value or "").replace("đ", "d").replace("Đ", "D")
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]", "", text.casefold())


def prepare_employee_report(employees: list[dict[str, Any]]) -> list[dict[str, Any]]:
    for employee in employees:
        registered = employee["invited"]
        try:
            target = float(str(employee["target"]).replace(",", "."))
        except (TypeError, ValueError):
            target = 0
        denominator = max(registered, target)
        employee["attendance_rate"] = (
            (employee["afternoon"] + employee["dinner"]) / denominator if denominator else None
        )
    return sorted(
        employees,
        key=lambda employee: (
            not bool(normalize(employee["department"])),
            normalize(employee["department"]),
            normalize(employee["name"]),
        ),
    )


def resolve_target_employee(inviter: Any, employees: list[dict[str, Any]]) -> dict[str, Any] | None:
    raw = normalize_header(inviter)
    if not raw or re.search(r"[,;/+&]", raw):
        return None
    if normalize(inviter) in {"tgd", "tdg", "septung"}:
        aliases = [
            employee for employee in employees
            if normalize(employee["name"]) == normalize("ANH TÙNG")
        ]
        if len(aliases) == 1:
            return aliases[0]
    exact = [employee for employee in employees if normalize(employee["name"]) == normalize(inviter)]
    if len(exact) == 1:
        return exact[0]

    words = raw.split()
    honorifics = {"anh", "chi", "chu", "co", "bac", "mr", "mrs", "sep"}
    words = [word for word in words if word not in honorifics]
    if not words or len(words) > 4:
        return None
    qualifiers = {"kt": {"ke", "toan"}, "kd": {"kinh", "doanh"}}
    qualifier_words = set().union(*(qualifiers[word] for word in words if word in qualifiers))
    name_words = {word for word in words if word not in qualifiers}
    if not name_words:
        name_words = set(words)
        qualifier_words = set()

    candidates = []
    for employee in employees:
        employee_name = set(normalize_header(employee["name"]).split())
        employee_details = set(
            normalize_header(f"{employee['position']} {employee['department']}").split()
        )
        matches_name = name_words <= employee_name
        matches_qualifiers = qualifier_words <= employee_details
        matches_position_only = not qualifier_words and name_words <= set(
            normalize_header(employee["position"]).split()
        )
        if (matches_name and matches_qualifiers) or matches_position_only:
            candidates.append(employee)
    return candidates[0] if len(candidates) == 1 else None


def summarize_unmatched_inviter_counts(
    charts: dict[str, list[dict[str, Any]]],
    employees: list[dict[str, Any]],
) -> dict[str, int]:
    chart_categories = {
        "registered": "registered",
        "afternoon": "afternoon",
        "dinner_invited": "dinner_invited",
        "dinner": "dinner",
    }
    return {
        category: sum(
            group["count"]
            for group in charts.get(category, [])
            if resolve_target_employee(group["name"], employees) is None
        )
        for category in chart_categories
    }


def normalize_header(value: Any) -> str:
    text = str(value or "").replace("đ", "d").replace("Đ", "D").casefold()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def column_index(headers: list[Any], field: str) -> int | None:
    normalized = [normalize_header(value) for value in headers]
    aliases = ALIASES[field]
    for alias in aliases:
        if alias in normalized:
            return normalized.index(alias)
    for index, header in enumerate(normalized):
        if field == "inviter" and re.search(r"\bsale\b", header):
            return index
        if any(alias in header for alias in aliases if len(alias) > 4):
            return index
    return None


def find_header(sheet: Any) -> tuple[int, list[Any]]:
    row_limit = min(sheet.max_row or 1, 10)
    preview = list(sheet.iter_rows(min_row=1, max_row=row_limit, values_only=True))
    header_index = 0
    best_score = -1
    for index, candidate in enumerate(preview):
        score = sum(
            column_index(list(candidate), field) is not None
            for field in ("employee", "target", "customer_id", "phone", "customer_name", "inviter", "vip", "guest_count", "checked_in")
        )
        if score > best_score:
            best_score = score
            header_index = index
    return header_index, list(preview[header_index]) if preview else []


def read_sheet(workbook: Any, logical_name: str) -> list[dict[str, str]]:
    sheet = workbook[logical_name]
    header_index, headers = find_header(sheet)
    if not any(clean(value) for value in headers):
        return []
    columns = {field: column_index(headers, field) for field in ALIASES}
    records = []
    for row in sheet.iter_rows(min_row=header_index + 2, values_only=True):
        values = list(row)
        if not any(clean(value) for value in values):
            continue
        record = {
            clean(header): clean(values[index]) if index is not None and index < len(values) else ""
            for index, header in enumerate(headers)
            if clean(header)
        }
        for field, index in columns.items():
            record[f"_{field}"] = clean(values[index]) if index is not None and index < len(values) else ""
        records.append(record)
    return records


def identity_key(customer: dict[str, Any]) -> str:
    for field in ("_customer_id", "_phone", "_customer_name"):
        value = normalize(customer.get(field))
        if value:
            return f"{field}:{value}"
    return ""


def guest_vip(customer: dict[str, Any]) -> bool:
    value = normalize(customer.get("_vip"))
    if value in ("khongvip", "kvip", "nonvip", "novip", "false", "no", "0"):
        return False
    raw_value = clean(customer.get("_vip")).casefold()
    return raw_value in ("ü", "✓", "✔", "x") or any(token in value for token in ("vip", "true", "yes", "co", "1"))


def is_checked_in(value: Any) -> bool:
    raw_value = clean(value).casefold()
    return raw_value in ("ü", "✓", "✔", "x", "yes", "true", "1", "co", "checked", "checked in")


def guest_count(row: dict[str, str]) -> int | None:
    raw_value = row.get("_guest_count", "")
    if not raw_value:
        return None
    try:
        numeric_count = float(raw_value.replace(",", "."))
    except (ValueError, TypeError):
        return None
    if not numeric_count.is_integer() or numeric_count <= 0:
        return None
    return int(numeric_count)


def customer_identity_value(row: dict[str, str], field: str) -> str:
    value = row.get(f"_{field}", "")
    if field == "phone":
        digits = re.sub(r"\D", "", value)
        if digits.startswith("84") and len(digits) >= 11:
            digits = "0" + digits[2:]
        return digits
    if field == "email":
        return value.strip().casefold()
    return normalize(value)


def link_customer_presence(
    source_rows: dict[str, tuple[str, list[dict[str, str]]]],
) -> dict[str, list[dict[str, bool]]]:
    categories = ("registered", "afternoon", "dinner")
    nodes = {
        (category, index)
        for category in categories
        for index in range(len(source_rows[category][1]))
    }
    parents = {node: node for node in nodes}
    member_categories = {node: {node[0]} for node in nodes}
    member_phones = {
        node: {
            phone
        } if (phone := customer_identity_value(source_rows[node[0]][1][node[1]], "phone")) else set()
        for node in nodes
    }
    member_emails = {
        node: {
            email
        } if (email := customer_identity_value(source_rows[node[0]][1][node[1]], "email")) else set()
        for node in nodes
    }

    def find(node: tuple[str, int]) -> tuple[str, int]:
        while parents[node] != node:
            parents[node] = parents[parents[node]]
            node = parents[node]
        return node

    def union(left: tuple[str, int], right: tuple[str, int], match_field: str) -> None:
        left_root = find(left)
        right_root = find(right)
        if left_root == right_root:
            return
        if member_categories[left_root] & member_categories[right_root]:
            return
        phones = member_phones[left_root] | member_phones[right_root]
        emails = member_emails[left_root] | member_emails[right_root]
        if match_field == "customer_name" and (
            (member_phones[left_root] and member_phones[right_root] and len(phones) > 1)
            or (member_emails[left_root] and member_emails[right_root] and len(emails) > 1)
        ):
            return
        parents[right_root] = left_root
        member_categories[left_root].update(member_categories.pop(right_root))
        member_phones[left_root] = phones
        member_phones.pop(right_root)
        member_emails[left_root] = emails
        member_emails.pop(right_root)

    for left_index, left_category in enumerate(categories):
        left_rows = source_rows[left_category][1]
        for right_category in categories[left_index + 1:]:
            right_rows = source_rows[right_category][1]
            for field in ("phone", "email", "customer_name"):
                left_values: dict[str, int] = {}
                right_values: dict[str, int] = {}
                left_duplicates: set[str] = set()
                right_duplicates: set[str] = set()
                for index, row in enumerate(left_rows):
                    value = customer_identity_value(row, field)
                    if value in left_values:
                        left_duplicates.add(value)
                    elif value:
                        left_values[value] = index
                for index, row in enumerate(right_rows):
                    value = customer_identity_value(row, field)
                    if value in right_values:
                        right_duplicates.add(value)
                    elif value:
                        right_values[value] = index
                for value in left_values.keys() & right_values.keys():
                    if value in left_duplicates or value in right_duplicates:
                        continue
                    union(
                        (left_category, left_values[value]),
                        (right_category, right_values[value]),
                        field,
                    )

    status_by_root: dict[tuple[str, int], dict[str, bool]] = {}
    for node in nodes:
        root = find(node)
        status = status_by_root.setdefault(
            root,
            {
                "registered": False,
                "afternoon": False,
                "dinner_invited": False,
                "dinner": False,
            },
        )
        category, index = node
        status[category if category != "dinner" else "dinner_invited"] = True
        if category == "dinner" and is_checked_in(
            source_rows[category][1][index].get("_checked_in", "")
        ):
            status["dinner"] = True

    return {
        category: [status_by_root[find((category, index))] for index in range(len(source_rows[category][1]))]
        for category in categories
    }


def classify_separate_file(filename: str, headers: list[str]) -> str:
    filename_key = normalize(filename)
    normalized_headers = {normalize_header(value) for value in headers}
    if "target" in filename_key or ("ten nhan vien" in normalized_headers and "so khach target" in normalized_headers):
        return "employees"
    if "chieu" in filename_key or "buoi chieu" in " ".join(normalized_headers):
        return "afternoon"
    if "tiectoi" in filename_key or ("sl khach" in normalized_headers and "check in" in normalized_headers):
        return "dinner"
    if "dangki" in filename_key or "dangky" in filename_key or (
        "bo phan moi" in normalized_headers and "nhan vien moi" in normalized_headers
    ):
        return "registered"
    if {"ho ten", "sdt", "nguoi moi", "vip"}.issubset(normalized_headers):
        return "afternoon"
    if {"ho ten", "sdt", "email", "thanh pho", "nhan vien moi"}.issubset(normalized_headers):
        return "registered"
    raise ValueError(f"Không xác định được loại dữ liệu của file “{filename}”. Hãy chọn đúng 4 file Excel đã cung cấp.")


def city_region(value: Any) -> str | None:
    normalized = normalize(value)
    if normalized in {"hn", "hanoi", "thanhphohanoi"}:
        return "HN"
    if normalized in {"hcm", "tphcm", "hochiminh", "thanhphohochiminh", "saigon"}:
        return "HCM"
    return None


def parse_separate_files(files: list[tuple[str, bytes]], region: str = "HN") -> dict[str, Any]:
    if region not in REGIONS:
        raise ValueError("Miền báo cáo không hợp lệ.")
    sources: dict[str, tuple[str, list[dict[str, str]]]] = {}
    region_warnings: list[str] = []
    missing_inviter_columns: set[str] = set()
    chart_group_by = {category: "Người mời" for category in ("registered", "afternoon", "dinner_invited", "dinner")}
    required_sources = {"registered", "afternoon", "dinner", "employees"}
    for filename, file_bytes in files:
        try:
            workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
        except Exception as error:
            raise ValueError(f"Không đọc được file Excel “{filename}”. Hãy chọn file .xlsx hợp lệ.") from error
        try:
            if not workbook.worksheets:
                raise ValueError(f"File “{filename}” không có trang tính dữ liệu.")
            sheet = workbook.worksheets[0]
            header_index, headers = find_header(sheet)
            role = classify_separate_file(filename, headers)
            required_fields = {
                "registered": ("customer_name", "phone", "employee"),
                "afternoon": ("customer_name", "inviter"),
                "dinner": ("customer_name", "guest_count", "checked_in"),
                "employees": ("employee", "target"),
            }[role]
            if any(column_index(headers, field) is None for field in required_fields):
                raise ValueError(f"File “{filename}” thiếu cột cần thiết để đọc đầy đủ dữ liệu.")
            if role in {"afternoon", "dinner"} and column_index(headers, "inviter") is None:
                missing_inviter_columns.add(role)
                if role == "dinner" and column_index(headers, "channel") is not None:
                    chart_group_by["dinner"] = "Kênh"
                    chart_group_by["dinner_invited"] = "Kênh"
            rows = read_sheet(workbook, sheet.title)
            if role in {"registered", "afternoon"} and column_index(headers, "city") is not None:
                kept_rows = []
                excluded_count = 0
                for row in rows:
                    row_region = city_region(row.get("_city", ""))
                    if row_region is not None and row_region != region:
                        excluded_count += 1
                    else:
                        kept_rows.append(row)
                rows = kept_rows
                if excluded_count:
                    region_warnings.append(
                        f"Đã bỏ qua {excluded_count} dòng thuộc miền khác trong file “{filename}”."
                    )
            if role == "dinner":
                rows = [
                    row for row in rows
                    if any(row.get(f"_{field}") for field in ("customer_name", "guest_count", "checked_in"))
                ]
        finally:
            workbook.close()
        if role in sources:
            raise ValueError(f"Đã chọn nhiều hơn một file cho mục “{role}”.")
        sources[role] = (filename, rows)

    missing = sorted(required_sources - sources.keys())
    if missing:
        labels = {
            "registered": "đăng ký tham dự",
            "afternoon": "khách check-in buổi chiều",
            "dinner": "khách mời tiệc tối",
            "employees": "chỉ tiêu nhân viên",
        }
        raise ValueError("Còn thiếu file: " + ", ".join(labels[source] for source in missing) + ".")

    employee_rows = sources["employees"][1]
    customer_presence = link_customer_presence(sources)
    employees_by_key: dict[str, dict[str, Any]] = {}
    for row in employee_rows:
        name = row.get("_employee", "")
        if not name:
            continue
        employees_by_key[normalize(name)] = {
            "name": name,
            "position": row.get("_position", ""),
            "department": row.get("_department", ""),
            "target": row.get("_target", ""),
            "invited": 0,
            "afternoon": 0,
            "dinner_invited": 0,
            "dinner": 0,
            "guests": [],
        }
    target_employee_count = len(employees_by_key)

    chart_groups: dict[str, dict[str, dict[str, Any]]] = {
        category: defaultdict(
            lambda: {"name": "", "vip": 0, "non_vip": 0, "unclassified": 0, "count": 0}
        )
        for category in ("registered", "afternoon", "dinner_invited", "dinner")
    }
    for source_category in ("registered", "afternoon", "dinner"):
        filename, rows = sources[source_category]
        for row_index, row in enumerate(rows):
            inviter = row.get("_employee", "") if source_category == "registered" else row.get("_inviter", "")
            inviter = inviter or "Chưa xác định"
            matched_employee = resolve_target_employee(inviter, list(employees_by_key.values()))
            employee_key = normalize(matched_employee["name"]) if matched_employee else ""
            quantity = guest_count(row) if source_category == "dinner" else 1
            measurable_quantity = quantity or 0
            checked_in = source_category == "dinner" and is_checked_in(row.get("_checked_in", ""))
            vip_value = row.get("_vip", "")
            vip_marked = guest_vip(row)
            vip_classified = bool(vip_value)
            employee = employees_by_key.get(employee_key) if employee_key else None
            if employee is not None:
                if source_category == "registered":
                    employee["invited"] += measurable_quantity
                elif source_category == "afternoon":
                    employee["afternoon"] += measurable_quantity
                elif source_category == "dinner":
                    employee["dinner_invited"] += measurable_quantity
                    if checked_in:
                        employee["dinner"] += measurable_quantity
                employee["guests"].append({
                    "customer_id": row.get("_customer_id", ""),
                    "name": row.get("_customer_name", ""),
                    "phone": row.get("_phone", ""),
                    **customer_presence[source_category][row_index],
                    "vip": vip_marked,
                    "vip_classified": vip_classified,
                    "quantity": quantity,
                    "source": source_category,
                    "source_filename": filename,
                    "source_data": {key: value for key, value in row.items() if not key.startswith("_")},
                })
            chart_categories = (
                ("registered",) if source_category == "registered"
                else ("afternoon",) if source_category == "afternoon"
                else ("dinner_invited", "dinner") if checked_in
                else ("dinner_invited",)
            )
            for chart_category in chart_categories:
                chart_inviter = matched_employee["name"] if matched_employee else inviter
                if source_category == "dinner" and chart_group_by[chart_category] == "Kênh":
                    chart_inviter = row.get("_channel", "") or "Chưa xác định"
                group = chart_groups[chart_category][normalize(chart_inviter)]
                group["name"] = chart_inviter
                group["count"] += measurable_quantity
                if measurable_quantity:
                    if vip_marked:
                        group["vip"] += measurable_quantity
                    elif vip_classified:
                        group["non_vip"] += measurable_quantity
                    else:
                        group["unclassified"] += measurable_quantity

    employees = prepare_employee_report(list(employees_by_key.values()))
    charts = {
        category: sorted(groups.values(), key=lambda item: (-item["count"], normalize(item["name"])))
        for category, groups in chart_groups.items()
    }
    unmatched = summarize_unmatched_inviter_counts(charts, employees)
    dinner_rows = sources["dinner"][1]
    dinner_invited = sum(guest_count(row) or 0 for row in dinner_rows)
    dinner_checked_in = sum(
        (guest_count(row) or 0) for row in dinner_rows if is_checked_in(row.get("_checked_in", ""))
    )
    dinner_unquantified = sum(guest_count(row) is None for row in dinner_rows)
    dinner_unquantified_checkins = sum(
        guest_count(row) is None and is_checked_in(row.get("_checked_in", ""))
        for row in dinner_rows
    )
    source_labels = {
        "registered": "đăng ký",
        "afternoon": "khách check-in buổi chiều",
        "dinner": "khách mời tiệc tối",
    }
    warnings = region_warnings + [
        f"{sum(not (row.get('_employee') if category == 'registered' else row.get('_inviter')) for row in sources[category][1])} dòng {source_labels[category]} chưa có tên người mời; các dòng này vẫn có trong nhóm “Chưa xác định”."
        for category in ("registered", "afternoon", "dinner")
        if category not in missing_inviter_columns
        if any(not (row.get("_employee") if category == "registered" else row.get("_inviter")) for row in sources[category][1])
    ]
    if "dinner" in missing_inviter_columns and dinner_invited:
        warnings.append(
            f"File “{sources['dinner'][0]}” không có cột người mời; {dinner_invited} khách tiệc tối vẫn được tính vào tổng nhưng nằm trong nhóm “Chưa xác định”."
        )
    if dinner_unquantified:
        warnings.append(
            f"{dinner_unquantified} dòng tiệc tối thiếu hoặc sai số lượng ở cột SL Khách; trong đó {dinner_unquantified_checkins} dòng đã đánh dấu check-in. Chưa tính các dòng này vào tổng khách/check-in."
        )
    return {
        "employees": employees,
        "charts": charts,
        "chart_group_by": chart_group_by,
        "unmatched": unmatched,
        "totals": {
            "employees": target_employee_count,
            "registered": len(sources["registered"][1]),
            "afternoon": len(sources["afternoon"][1]),
            "dinner_invited": dinner_invited,
            "dinner": dinner_checked_in,
            "dinner_unquantified": dinner_unquantified,
            "dinner_unquantified_checkins": dinner_unquantified_checkins,
        },
        "warnings": warnings,
        "sources": [
            {
                "category": category,
                "filename": filename,
                "records": len(rows),
                "guests": sum(guest_count(row) or 0 for row in rows) if category == "dinner" else len(rows),
                "unquantified_records": sum(guest_count(row) is None for row in rows) if category == "dinner" else 0,
                "unquantified_checkins": sum(
                    guest_count(row) is None and is_checked_in(row.get("_checked_in", ""))
                    for row in rows
                ) if category == "dinner" else 0,
            }
            for category, (filename, rows) in sources.items()
        ],
    }


def parse_workbook(file_bytes: bytes) -> dict[str, Any]:
    try:
        workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as error:
        raise ValueError("Không đọc được file Excel. Hãy tải lên file .xlsx hợp lệ.") from error

    actual_by_normalized = {normalize(name): name for name in workbook.sheetnames}
    missing = [
        required
        for normalized, required in (
            ("dskhcheckintiectoi", "Ds_kh_checkin_tiectoi"),
            ("dskhcheckinchieu", "Ds_kh_checkin_chieu"),
            ("dskhdkqr", "Ds_kh_dk_qr"),
            ("dstargetnv", "Ds_target_nv"),
        )
        if normalized not in actual_by_normalized
    ]
    if missing:
        raise ValueError("Thiếu sheet bắt buộc: " + ", ".join(missing))

    sheet_names = {key: actual_by_normalized[key] for key in REQUIRED_SHEETS}
    sources = {
        key: read_sheet(workbook, sheet_names[key])
        for key in ("dskhcheckintiectoi", "dskhcheckinchieu", "dskhdkqr")
    }
    target_rows = read_sheet(workbook, sheet_names["dstargetnv"])
    workbook.close()

    employees_by_key: dict[str, dict[str, Any]] = {}
    for row in target_rows:
        name = row.get("_employee", "")
        if not name:
            continue
        key = normalize(name)
        employees_by_key[key] = {
            "name": name,
            "position": row.get("_position", ""),
            "department": row.get("_department", ""),
            "target": row.get("_target", ""),
            "customers": {},
        }

    customer_rows: dict[str, dict[str, dict[str, Any]]] = {
        category: {} for category in ("registered", "afternoon", "dinner")
    }
    all_customers: dict[str, dict[str, Any]] = {}

    for category, source_key in (
        ("registered", "dskhdkqr"),
        ("afternoon", "dskhcheckinchieu"),
        ("dinner", "dskhcheckintiectoi"),
    ):
        for row in sources[source_key]:
            inviter = row.get("_inviter", "")
            employee_key = normalize(inviter)
            key = identity_key(row)
            if not key:
                key = f"row:{category}:{len(all_customers)}"
            customer = all_customers.setdefault(
                key,
                {
                    "key": key,
                    "customer_id": row.get("_customer_id", ""),
                    "name": row.get("_customer_name", ""),
                    "phone": row.get("_phone", ""),
                    "inviter": inviter,
                    "vip": False,
                    "registered": False,
                    "afternoon": False,
                    "dinner": False,
                },
            )
            customer["customer_id"] = customer["customer_id"] or row.get("_customer_id", "")
            customer["name"] = customer["name"] or row.get("_customer_name", "")
            customer["phone"] = customer["phone"] or row.get("_phone", "")
            customer["inviter"] = customer["inviter"] or inviter
            customer["vip"] = customer["vip"] or guest_vip(row)
            customer[category] = True
            customer_rows[category][key] = customer
            if employee_key:
                employee = employees_by_key.setdefault(
                    employee_key,
                    {"name": inviter, "position": "", "department": "", "target": "", "customers": {}},
                )
                employee["customers"][key] = customer
                if not customer["inviter"]:
                    customer["inviter"] = employee["name"]

    employees = []
    for employee in employees_by_key.values():
        guests = list(employee.pop("customers").values())
        employees.append(
            {
                "name": employee["name"],
                "position": employee["position"],
                "department": employee["department"],
                "target": employee["target"],
                "invited": sum(guest["registered"] for guest in guests),
                "afternoon": sum(guest["afternoon"] for guest in guests),
                "dinner": sum(guest["dinner"] for guest in guests),
                "guests": guests,
            }
        )
    employees = prepare_employee_report(employees)

    charts = {}
    for category in ("dinner", "afternoon", "registered"):
        groups: dict[str, dict[str, Any]] = defaultdict(lambda: {"name": "", "vip": 0, "non_vip": 0, "count": 0})
        for customer in customer_rows[category].values():
            inviter = customer["inviter"] or "Chưa xác định"
            item = groups[normalize(inviter)]
            item["name"] = inviter
            item["count"] += 1
            item["vip" if customer["vip"] else "non_vip"] += 1
        charts[category] = sorted(
            groups.values(),
            key=lambda item: (-item["count"], normalize(item["name"])),
        )

    unmatched = summarize_unmatched_inviter_counts(
        {**charts, "dinner_invited": charts["dinner"]},
        employees,
    )
    target_employee_rows = list(employees_by_key.values())
    dinner_rows = sources["dskhcheckintiectoi"]
    unmatched["dinner_invited"] = sum(
        guest_count(row) or 0
        for row in dinner_rows
        if resolve_target_employee(row.get("_inviter", ""), target_employee_rows) is None
    )
    unmatched["dinner"] = sum(
        (guest_count(row) or 0)
        for row in dinner_rows
        if resolve_target_employee(row.get("_inviter", ""), target_employee_rows) is None
        and is_checked_in(row.get("_checked_in", ""))
    )
    return {
        "employees": employees,
        "charts": charts,
        "unmatched": unmatched,
        "totals": {
            "employees": len(employees),
            "registered": len(customer_rows["registered"]),
            "afternoon": len(customer_rows["afternoon"]),
            "dinner": len(customer_rows["dinner"]),
        },
        "sheets": {category: sheet_names[key] for category, key in (
            ("dinner", "dskhcheckintiectoi"),
            ("afternoon", "dskhcheckinchieu"),
            ("registered", "dskhdkqr"),
            ("employees", "dstargetnv"),
        )},
    }


@app.get("/")
def index():
    return render_template("index.html", auth_required=AUTH_REQUIRED)


@app.get("/api/report")
def get_report():
    if DATABASE_URL:
        try:
            reports = reports_from_payload(read_persisted_report())
        except psycopg.Error:
            app.logger.exception("Không thể đọc báo cáo đã lưu trên Supabase.")
            return jsonify({"ok": False, "message": "Không thể tải báo cáo đã lưu. Vui lòng thử lại sau."}), 503
    else:
        with _report_lock:
            reports = dict(_current_reports)
    if not reports:
        return jsonify({"ok": False, "message": "Chưa có báo cáo. Hãy tải lên 4 file Excel."}), 404
    region = request.args.get("region", "HN").upper()
    if region not in REGIONS:
        return jsonify({"ok": False, "message": "Miền báo cáo không hợp lệ."}), 400
    return jsonify({
        "ok": True,
        "report": reports.get(region),
        "reports": reports,
        "region": region,
    })


@app.post("/api/upload")
def upload():
    uploaded = request.files.get("file")
    uploaded_files = request.files.getlist("files")
    if uploaded is None and not uploaded_files:
        return jsonify({"ok": False, "message": "Vui lòng chọn 4 file Excel .xlsx."}), 400
    try:
        region = request.form.get("region", "HN").strip().upper()
        if region not in REGIONS:
            raise ValueError("Vui lòng chọn miền Hà Nội hoặc TP. Hồ Chí Minh.")
        if uploaded is not None:
            if not uploaded.filename or not uploaded.filename.lower().endswith(".xlsx"):
                raise ValueError("Chỉ hỗ trợ file Excel định dạng .xlsx.")
            report = parse_workbook(uploaded.read())
        else:
            if any(not file.filename or not file.filename.lower().endswith(".xlsx") for file in uploaded_files):
                raise ValueError("Chỉ hỗ trợ file Excel định dạng .xlsx.")
            if len(uploaded_files) == 1:
                report = parse_workbook(uploaded_files[0].read())
            else:
                report = parse_separate_files([
                    (file.filename, file.read())
                    for file in uploaded_files
                ], region=region)
    except ValueError as error:
        return jsonify({"ok": False, "message": str(error)}), 400
    if DATABASE_URL:
        try:
            reports = reports_from_payload(read_persisted_report())
            reports[region] = report
            write_persisted_report({"regions": reports})
        except psycopg.Error:
            app.logger.exception("Không thể cập nhật báo cáo trên Supabase.")
            return jsonify({"ok": False, "message": "Không thể lưu báo cáo lâu dài. Vui lòng thử lại sau."}), 503
    with _report_lock:
        if not DATABASE_URL:
            reports = dict(_current_reports)
            reports[region] = report
        _current_reports.clear()
        _current_reports.update(reports)
    return jsonify({
        "ok": True,
        "report": report,
        "reports": reports,
        "region": region,
    })


@app.post("/api/report/adjust")
def adjust_report():
    payload = request.get_json(silent=True)
    if not isinstance(payload, dict):
        return jsonify({"ok": False, "message": "Dữ liệu điều chỉnh không hợp lệ."}), 400
    region = str(payload.get("region", "")).strip().upper()
    if region != "HCM":
        return jsonify({"ok": False, "message": "Chỉ hỗ trợ điều chỉnh check-in tiệc tối cho HCM."}), 400
    total = payload.get("dinner_checkins")
    if isinstance(total, bool) or not isinstance(total, int):
        return jsonify({"ok": False, "message": "Số check-in tiệc tối phải là số nguyên."}), 400
    reason = str(payload.get("reason", "")).strip()
    if not reason:
        return jsonify({"ok": False, "message": "Vui lòng ghi rõ lý do điều chỉnh."}), 400

    if DATABASE_URL:
        try:
            reports = reports_from_payload(read_persisted_report())
            report = reports.get(region)
            if report is None:
                return jsonify({"ok": False, "message": "Chưa có báo cáo HCM để điều chỉnh."}), 404
            adjusted = set_dinner_checkin_total(report, total, reason)
            reports[region] = adjusted
            write_persisted_report({"regions": reports})
        except ValueError as error:
            return jsonify({"ok": False, "message": str(error)}), 400
        except psycopg.Error:
            app.logger.exception("Không thể lưu điều chỉnh báo cáo lên Supabase.")
            return jsonify({"ok": False, "message": "Không thể lưu điều chỉnh lâu dài. Vui lòng thử lại sau."}), 503
    else:
        with _report_lock:
            reports = dict(_current_reports)
            report = reports.get(region)
            if report is None:
                return jsonify({"ok": False, "message": "Chưa có báo cáo HCM để điều chỉnh."}), 404
            try:
                adjusted = set_dinner_checkin_total(report, total, reason)
            except ValueError as error:
                return jsonify({"ok": False, "message": str(error)}), 400
            reports[region] = adjusted
            _current_reports.clear()
            _current_reports.update(reports)
    return jsonify({"ok": True, "region": region, "report": adjusted, "reports": reports})


@app.errorhandler(RequestEntityTooLarge)
def upload_too_large(_error):
    del _error
    return jsonify({"ok": False, "message": "File vượt quá giới hạn 20 MB."}), 413


def load_default_report() -> None:
    source_dir = Path(environ.get("REPORT_DATA_DIR", Path.home() / "Downloads" / "15.10"))
    if not source_dir.is_dir():
        return
    files = sorted(
        file for file in source_dir.glob("*.xlsx")
        if not file.name.startswith("~$")
    )
    if not files:
        return
    try:
        report = parse_separate_files([(file.name, file.read_bytes()) for file in files])
    except (OSError, ValueError) as error:
        app.logger.error("Không thể tự nạp báo cáo từ %s: %s", source_dir, error)
        return
    global _current_reports
    with _report_lock:
        _current_reports = {"HN": report}
    app.logger.info("Đã nạp báo cáo từ %s", source_dir)


if __name__ == "__main__":
    load_default_report()
    app.run(debug=False)
